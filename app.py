import os
import io
import re
from datetime import datetime, timedelta
from functools import wraps

from flask import (Flask, render_template, request, redirect, url_for,
                   session, send_file, flash)
from werkzeug.utils import secure_filename
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from forms import LoginForm, UserRegisterForm  # noqa: F401
from services.database import Database

app = Flask(__name__)
# En production : définir la variable d'environnement SECRET_KEY (ne jamais garder la valeur par défaut)
app.secret_key = os.environ.get('SECRET_KEY', 'une_cle_secrete_tres_robuste_par_defaut')
app.permanent_session_lifetime = timedelta(minutes=30)

db = Database()

TAUX_USD_DEFAUT = 1.08
EXTENSIONS_LOGO = {'png', 'jpg', 'jpeg', 'webp', 'svg'}


@app.before_request
def make_session_permanent():
    session.permanent = True


def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function


def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session or session.get('role') != 'admin':
            flash('Accès réservé aux administrateurs.', 'error')
            return redirect(url_for('dashboard'))
        return f(*args, **kwargs)
    return decorated_function


# ----------------------------------------------------------------------
# TRADUCTIONS (FR / EN)
# ----------------------------------------------------------------------
TRADUCTIONS = {
    'fr': {
        'dashboard': 'Tableau de bord', 'ventes': 'Ventes', 'produits': 'Produits & Stock',
        'clients': 'Clients', 'depenses': 'Dépenses', 'rapports': 'Rapports',
        'parametres': 'Paramètres', 'logout': 'Déconnexion',
        'entreprise': "Gestion d'entreprise", 'menu': 'Menu de navigation', 'connecté': 'En ligne'
    },
    'en': {
        'dashboard': 'Dashboard', 'ventes': 'Sales', 'produits': 'Products & Stock',
        'clients': 'Clients', 'depenses': 'Expenses', 'rapports': 'Reports',
        'parametres': 'Settings', 'logout': 'Logout',
        'entreprise': 'Business Management', 'menu': 'Navigation Menu', 'connecté': 'Online'
    }
}


@app.route('/changer-langue/<lang>', methods=['POST'])
@login_required
def changer_langue(lang):
    if lang in ['fr', 'en']:
        session['lang'] = lang
    return redirect(request.referrer or url_for('dashboard'))


# ----------------------------------------------------------------------
# VALIDATION
# ----------------------------------------------------------------------
REGEX_EMAIL = r'^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$'
# Madagascar : +261 ou 0, puis 32 / 33 / 34 / 37 / 38, puis 7 chiffres
REGEX_PHONE_MG = r'^(\+261|0)(32|33|34|37|38)\d{7}$'


def valider_email(email):
    if not email:
        return True
    return bool(re.match(REGEX_EMAIL, email.strip()))


def valider_telephone(phone):
    if not phone:
        return True
    return bool(re.match(REGEX_PHONE_MG, phone.strip().replace(" ", "")))


# ----------------------------------------------------------------------
# DEVISES : les montants sont stockés en EUROS et convertis à l'affichage
# ----------------------------------------------------------------------
def lire_parametres():
    with db.get_connection() as conn:
        return conn.execute('SELECT * FROM parametres LIMIT 1').fetchone()


def convertir_montant(valeur_eur, devise, taux_ar):
    if devise == 'Ar':
        return valeur_eur * taux_ar
    if devise == '$':
        return valeur_eur * TAUX_USD_DEFAUT
    return valeur_eur


def formater_montant(valeur_eur, devise, taux_ar):
    v = convertir_montant(valeur_eur, devise, taux_ar)
    if devise == 'Ar':
        return f"{round(v):,}".replace(',', ' ') + ' Ar'
    return f"{v:,.2f}".replace(',', ' ') + f' {devise}'


@app.context_processor
def inject_globals():
    with db.get_connection() as conn:
        param = conn.execute('SELECT * FROM parametres LIMIT 1').fetchone()
        # Cloche d'alertes : ruptures et stocks sous le seuil, les plus urgents d'abord
        alertes = conn.execute(
            "SELECT nom, stock, seuil_alerte FROM produits WHERE stock <= seuil_alerte "
            "ORDER BY stock ASC, nom ASC"
        ).fetchall()
    cles = param.keys() if param else []

    devise = session.get('devise', param['devise_defaut'] if 'devise_defaut' in cles else 'Ar')
    lang = session.get('lang', 'fr')
    accent_color = session.get('accent_color', param['couleur_accent'] if 'couleur_accent' in cles else '#8B5CF6')

    def translate(key):
        return TRADUCTIONS.get(lang, TRADUCTIONS['fr']).get(key, key)

    return {
        'current_devise': devise,
        'current_theme': session.get('theme', 'dark'),
        'current_lang': lang,
        'accent_color': accent_color,
        'nom_entreprise': param['nom_entreprise'] if 'nom_entreprise' in cles else 'Mini-ERP Pro',
        'tva_taux': param['tva'] if 'tva' in cles else 20.0,
        'taux_ar': (param['taux_ar'] if 'taux_ar' in cles and param['taux_ar'] else 5021.0),
        'taux_usd': TAUX_USD_DEFAUT,
        'logo_entreprise': param['logo'] if 'logo' in cles else None,
        'alertes_stock': alertes[:6],
        'alertes_stock_count': len(alertes),
        't': translate
    }


@app.route('/changer-devise/<devise>', methods=['POST'])
@login_required
def changer_devise(devise):
    if devise in ['Ar', '€', '$']:
        db.set_devise(session['user_id'], devise)
        session['devise'] = devise
    return redirect(request.referrer or url_for('dashboard'))


# ----------------------------------------------------------------------
# AUTHENTIFICATION
# ----------------------------------------------------------------------
@app.route('/login', methods=['GET', 'POST'])
def login():
    form = LoginForm()
    if 'failed_attempts' not in session:
        session['failed_attempts'] = 0
    if form.validate_on_submit():
        if session['failed_attempts'] >= 5:
            flash('Trop de tentatives infructueuses. Compte temporairement bloqué.', 'error')
            return render_template('login.html', form=form)
        user = db.verifier_utilisateur(form.username.data, form.password.data)
        if user:
            session.pop('failed_attempts', None)
            session['user_id'] = user['id']
            session['username'] = user['username']
            session['role'] = user['role']
            flash('Connexion réussie !', 'success')
            return redirect(url_for('dashboard'))
        session['failed_attempts'] += 1
        restantes = 5 - session['failed_attempts']
        if restantes > 0:
            flash(f'Identifiants incorrects. Il vous reste {restantes} tentative(s).', 'error')
        else:
            flash('Trop de tentatives échouées. Accès temporairement suspendu.', 'error')
    return render_template('login.html', form=form)


@app.route('/logout')
def logout():
    session.clear()
    flash('Déconnexion réussie.', 'success')
    return redirect(url_for('login'))


# ----------------------------------------------------------------------
# TABLEAU DE BORD : uniquement des données réelles (plus de valeurs inventées)
# ----------------------------------------------------------------------
@app.route('/')
@app.route('/dashboard')
@login_required
def dashboard():
    with db.get_connection() as conn:
        ca_total = conn.execute("SELECT SUM(total) FROM ventes").fetchone()[0] or 0.0
        depenses_total = conn.execute("SELECT SUM(montant) FROM depenses").fetchone()[0] or 0.0
        ventes_count = conn.execute("SELECT COUNT(*) FROM ventes").fetchone()[0] or 0
        # Alertes = ruptures + stocks sous le seuil
        stock_critique_count = conn.execute("SELECT COUNT(*) FROM produits WHERE stock = 0").fetchone()[0] or 0
        stock_alertes = conn.execute("SELECT COUNT(*) FROM produits WHERE stock <= seuil_alerte").fetchone()[0] or 0

    stats = {
        'ca': ca_total,
        'benefice': ca_total - depenses_total,
        'depenses': depenses_total,
        'ventes_count': ventes_count,
        'stock_bas_count': stock_alertes,
        'stock_critique_count': stock_critique_count,
    }

    rapport = db.recuperer_rapport_mensuel(6)  # 6 derniers mois réels, mois vides = 0
    chart_ca = {'labels': [r['mois'] for r in rapport], 'data': [r['ca'] for r in rapport]}

    depenses_cat = db.recuperer_depenses_par_categorie()
    chart_depenses = {
        'labels': [d['categorie'].capitalize() for d in depenses_cat],
        'data': [d['total'] for d in depenses_cat]
    }

    return render_template('dashboard.html',
                           stats=stats,
                           chart_ca=chart_ca,
                           chart_depenses=chart_depenses,
                           ventes_recentes=db.recuperer_ventes_recentes(5),
                           produits_reappro=db.recuperer_produits_a_reapprovisionner()[:5])


# ----------------------------------------------------------------------
# PARAMÈTRES (nom, devise, TVA, taux de change, couleur, logo)
# ----------------------------------------------------------------------
@app.route('/parametres', methods=['GET', 'POST'])
@login_required
@admin_required
def parametres():
    conn = db.get_connection()
    if request.method == 'POST':
        nom_entreprise = request.form.get('nom_entreprise', 'Mini-ERP Pro').strip() or 'Mini-ERP Pro'
        devise_defaut = request.form.get('devise_defaut', 'Ar')
        if devise_defaut not in ['Ar', '€', '$']:
            devise_defaut = 'Ar'

        try:
            tva = float(request.form.get('tva', 20.0))
            taux_ar = float(request.form.get('taux_ar', 5021.0))
            if tva < 0 or taux_ar <= 0:
                raise ValueError()
        except ValueError:
            flash("Valeurs numériques invalides dans les paramètres.", "error")
            conn.close()
            return redirect(url_for('parametres'))

        couleur_accent = request.form.get('couleur_accent', '#8B5CF6')
        if not re.match(r'^#[0-9a-fA-F]{6}$', couleur_accent):
            couleur_accent = '#8B5CF6'

        logo_file = request.files.get('logo')
        if logo_file and logo_file.filename:
            nom_fichier = secure_filename(logo_file.filename)
            extension = nom_fichier.rsplit('.', 1)[-1].lower() if '.' in nom_fichier else ''
            if extension not in EXTENSIONS_LOGO:
                flash("Format de logo non accepté (PNG, JPG, WEBP ou SVG).", "error")
                conn.close()
                return redirect(url_for('parametres'))
            os.makedirs('static/uploads', exist_ok=True)
            logo_path = f"static/uploads/{nom_fichier}"
            logo_file.save(logo_path)
            conn.execute('UPDATE parametres SET logo = ? WHERE id = 1', (logo_path,))

        conn.execute('''
            UPDATE parametres
            SET nom_entreprise = ?, devise_defaut = ?, tva = ?, couleur_accent = ?, taux_ar = ?
            WHERE id = 1
        ''', (nom_entreprise, devise_defaut, tva, couleur_accent, taux_ar))
        conn.commit()
        conn.close()
        session['devise'] = devise_defaut
        session['accent_color'] = couleur_accent
        flash('Paramètres mis à jour avec succès.', 'success')
        return redirect(url_for('parametres'))

    param = conn.execute('SELECT * FROM parametres WHERE id = 1').fetchone()
    conn.close()
    return render_template('parametres.html', param=param)


# ----------------------------------------------------------------------
# CLIENTS
# ----------------------------------------------------------------------
@app.route('/clients')
@login_required
def clients():
    recherche = request.args.get('q', '')
    page = max(int(request.args.get('page', 1) or 1), 1)
    per_page = 8
    search_term = f"%{recherche}%"

    with db.get_connection() as conn:
        total_clients = conn.execute(
            "SELECT COUNT(*) FROM clients WHERE nom LIKE ? OR email LIKE ? OR telephone LIKE ?",
            (search_term, search_term, search_term)).fetchone()[0]
        clients_list = conn.execute(
            "SELECT * FROM clients WHERE nom LIKE ? OR email LIKE ? OR telephone LIKE ? "
            "ORDER BY nom ASC LIMIT ? OFFSET ?",
            (search_term, search_term, search_term, per_page, (page - 1) * per_page)).fetchall()

    total_pages = (total_clients + per_page - 1) // per_page
    return render_template('clients.html', clients=clients_list, recherche=recherche,
                           page=page, total_pages=total_pages)


@app.route('/clients/ajouter', methods=['POST'])
@login_required
def ajouter_client():
    nom = request.form.get('nom', '').strip()
    email = request.form.get('email', '').strip()
    telephone = request.form.get('telephone', '').strip()
    if not nom:
        flash('Le nom du client est obligatoire.', 'error')
        return redirect(url_for('clients'))
    if email and not valider_email(email):
        flash("Format d'email invalide.", 'error')
        return redirect(url_for('clients'))
    if telephone and not valider_telephone(telephone):
        flash("Numéro de téléphone invalide (ex : +261 34 12 345 67).", 'error')
        return redirect(url_for('clients'))
    db.ajouter_client(nom, email, telephone)
    flash('Client ajouté avec succès.', 'success')
    return redirect(url_for('clients'))


@app.route('/clients/modifier/<int:client_id>', methods=['GET', 'POST'])
@login_required
def modifier_client(client_id):
    with db.get_connection() as conn:
        client = conn.execute("SELECT * FROM clients WHERE id = ?", (client_id,)).fetchone()
    if not client:
        flash('Client introuvable.', 'error')
        return redirect(url_for('clients'))
    if request.method == 'POST':
        nom = request.form.get('nom', '').strip()
        email = request.form.get('email', '').strip()
        telephone = request.form.get('telephone', '').strip()
        if not nom:
            flash('Le nom est obligatoire.', 'error')
        elif email and not valider_email(email):
            flash("Format d'email invalide.", 'error')
        elif telephone and not valider_telephone(telephone):
            flash("Numéro de téléphone invalide (ex : +261 34 12 345 67).", 'error')
        else:
            db.modifier_client(client_id, nom, email, telephone)
            flash('Client modifié avec succès.', 'success')
            return redirect(url_for('clients'))
    return render_template('modifier_client.html', client=client)


@app.route('/clients/supprimer/<int:client_id>', methods=['POST'])
@login_required
@admin_required
def supprimer_client(client_id):
    db.supprimer_client(client_id)
    flash('Client supprimé avec succès.', 'success')
    return redirect(url_for('clients'))


# ----------------------------------------------------------------------
# PRODUITS (pagination, tri, recherche)
# ----------------------------------------------------------------------
@app.route('/produits')
@login_required
def produits():
    recherche = request.args.get('q', '')
    tri = request.args.get('tri', 'nom')
    order = request.args.get('order', 'asc')
    page = max(int(request.args.get('page', 1) or 1), 1)
    per_page = 8

    col_tri = tri if tri in ['nom', 'prix', 'stock'] else 'nom'  # liste blanche : pas d'injection SQL
    dir_order = 'ASC' if order == 'asc' else 'DESC'
    search_term = f"%{recherche}%"

    with db.get_connection() as conn:
        total = conn.execute("SELECT COUNT(*) FROM produits WHERE nom LIKE ?", (search_term,)).fetchone()[0]
        produits_list = conn.execute(
            f"SELECT * FROM produits WHERE nom LIKE ? ORDER BY {col_tri} {dir_order} LIMIT ? OFFSET ?",
            (search_term, per_page, (page - 1) * per_page)).fetchall()

    total_pages = (total + per_page - 1) // per_page
    return render_template('produits.html', produits=produits_list, recherche=recherche,
                           page=page, total_pages=total_pages, tri=tri, order=order)


@app.route('/produits/ajouter', methods=['POST'])
@login_required
def ajouter_produit():
    nom = request.form.get('nom', '').strip()
    try:
        prix = float(request.form.get('prix', 0))
        stock = int(request.form.get('stock', 0))
        seuil_alerte = int(request.form.get('seuil_alerte', 5))
    except ValueError:
        flash('Montants ou stocks invalides.', 'error')
        return redirect(url_for('produits'))
    if not nom or prix <= 0 or stock < 0 or seuil_alerte < 0:
        flash('Veuillez saisir des valeurs positives valides.', 'error')
        return redirect(url_for('produits'))
    db.ajouter_produit(nom, prix, stock, seuil_alerte)
    flash('Produit ajouté avec succès.', 'success')
    return redirect(url_for('produits'))


@app.route('/produits/modifier/<int:produit_id>', methods=['GET', 'POST'])
@login_required
def modifier_produit(produit_id):
    with db.get_connection() as conn:
        produit = conn.execute("SELECT * FROM produits WHERE id = ?", (produit_id,)).fetchone()
    if not produit:
        flash('Produit introuvable.', 'error')
        return redirect(url_for('produits'))
    if request.method == 'POST':
        nom = request.form.get('nom', '').strip()
        try:
            prix = float(request.form.get('prix', 0))
            stock = int(request.form.get('stock', 0))
            seuil_alerte = int(request.form.get('seuil_alerte', 5))
        except ValueError:
            flash('Valeurs invalides.', 'error')
            return redirect(url_for('produits'))
        if nom and prix > 0 and stock >= 0 and seuil_alerte >= 0:
            db.modifier_produit(produit_id, nom, prix, stock, seuil_alerte)
            flash('Produit modifié avec succès.', 'success')
            return redirect(url_for('produits'))
        flash('Veuillez remplir tous les champs correctement.', 'error')
    return render_template('modifier_produit.html', produit=produit)


@app.route('/produits/supprimer/<int:produit_id>', methods=['POST'])
@login_required
@admin_required
def supprimer_produit(produit_id):
    db.supprimer_produit(produit_id)
    flash('Produit supprimé avec succès.', 'success')
    return redirect(url_for('produits'))


# ----------------------------------------------------------------------
# VENTES & CONTRÔLE DU STOCK
# ----------------------------------------------------------------------
@app.route('/ventes/nouvelle', methods=['GET', 'POST'])
@login_required
def nouvelle_vente():
    if request.method == 'POST':
        client_id = request.form.get('client_id')
        produit_ids = request.form.getlist('produit_id[]')
        quantites = request.form.getlist('quantite[]')

        # On regroupe les quantités par produit (même produit sur 2 lignes = 1 seul contrôle de stock)
        demandes = {}
        try:
            for p_id, q in zip(produit_ids, quantites):
                if p_id and q:
                    q_val = int(q)
                    if q_val <= 0:
                        flash('La quantité doit être supérieure à 0.', 'error')
                        return redirect(url_for('nouvelle_vente'))
                    demandes[int(p_id)] = demandes.get(int(p_id), 0) + q_val
        except ValueError:
            flash('Quantité invalide.', 'error')
            return redirect(url_for('nouvelle_vente'))

        if not demandes:
            flash('Veuillez sélectionner au moins un produit.', 'error')
            return redirect(url_for('nouvelle_vente'))

        lignes = []
        with db.get_connection() as conn:
            for p_id, q_val in demandes.items():
                prod = conn.execute("SELECT * FROM produits WHERE id = ?", (p_id,)).fetchone()
                if not prod:
                    continue
                if prod['stock'] < q_val:
                    flash(f"Stock insuffisant pour '{prod['nom']}' (stock disponible : {prod['stock']}).", 'error')
                    return redirect(url_for('nouvelle_vente'))
                lignes.append({'produit_id': p_id, 'quantite': q_val, 'prix_unitaire': prod['prix']})

        if not lignes:
            flash('Aucun produit valide sélectionné.', 'error')
            return redirect(url_for('nouvelle_vente'))
        try:
            db.creer_vente(int(client_id) if client_id else None, lignes)
        except ValueError as e:
            flash(str(e), 'error')
            return redirect(url_for('nouvelle_vente'))
        flash('Vente enregistrée avec succès !', 'success')
        return redirect(url_for('ventes'))

    return render_template('nouvelle_vente.html',
                           clients=db.recuperer_clients(), produits=db.recuperer_produits())


def charger_vente(vente_id):
    with db.get_connection() as conn:
        vente = conn.execute('''
            SELECT v.*, c.nom AS client_nom, c.email AS client_email, c.telephone AS client_telephone
            FROM ventes v LEFT JOIN clients c ON v.client_id = c.id
            WHERE v.id = ?
        ''', (vente_id,)).fetchone()
        lignes = conn.execute('''
            SELECT lv.*, p.nom AS produit_nom
            FROM lignes_vente lv LEFT JOIN produits p ON lv.produit_id = p.id
            WHERE lv.vente_id = ?
        ''', (vente_id,)).fetchall()
    return vente, lignes


@app.route('/ventes/<int:vente_id>')
@login_required
def detail_vente(vente_id):
    vente, lignes = charger_vente(vente_id)
    if not vente:
        flash('Vente introuvable.', 'error')
        return redirect(url_for('ventes'))
    return render_template('detail_vente.html', vente=vente, lignes=lignes)


@app.route('/ventes/pdf/<int:vente_id>')
@login_required
def telecharger_facture_pdf(vente_id):
    vente, lignes = charger_vente(vente_id)
    if not vente:
        flash('Vente introuvable.', 'error')
        return redirect(url_for('ventes'))

    param = lire_parametres()
    cles = param.keys() if param else []
    nom_ent = param['nom_entreprise'] if 'nom_entreprise' in cles else 'FACTURE'
    taux_ar = param['taux_ar'] if 'taux_ar' in cles and param['taux_ar'] else 5021.0
    devise = session.get('devise', param['devise_defaut'] if 'devise_defaut' in cles else 'Ar')

    def fm(valeur):
        return formater_montant(valeur, devise, taux_ar)

    buffer = io.BytesIO()
    p = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter

    # En-tête : nom de l'entreprise + logo (si présent et lisible)
    p.setFont("Helvetica-Bold", 18)
    p.drawString(50, height - 50, nom_ent)
    logo = param['logo'] if 'logo' in cles else None
    if logo and os.path.exists(logo) and not logo.lower().endswith('.svg'):
        try:
            p.drawImage(logo, width - 160, height - 85, width=110, height=55,
                        preserveAspectRatio=True, mask='auto')
        except Exception:
            pass  # un logo illisible ne doit jamais empêcher la facture

    p.setFont("Helvetica", 10)
    p.drawString(50, height - 75, f"Facture N° : #{vente['id']}")
    p.drawString(50, height - 90, f"Date : {vente['date']}")
    p.drawString(50, height - 115, f"Client : {vente['client_nom'] or 'Client comptoir'}")

    def entete_tableau(y):
        p.setFont("Helvetica-Bold", 10)
        p.drawString(50, y, "Article")
        p.drawRightString(340, y, "Qté")
        p.drawRightString(450, y, "P.U.")
        p.drawRightString(550, y, "Total")
        p.setLineWidth(0.5)
        p.line(50, y - 5, 550, y - 5)
        p.setFont("Helvetica", 10)
        return y - 22

    y = entete_tableau(height - 160)
    for ligne in lignes:
        if y < 90:  # nouvelle page si la facture est longue
            p.showPage()
            y = entete_tableau(height - 60)
        p.drawString(50, y, str(ligne['produit_nom'] or 'Article supprimé')[:48])
        p.drawRightString(340, y, str(ligne['quantite']))
        p.drawRightString(450, y, fm(ligne['prix_unitaire']))
        p.drawRightString(550, y, fm(ligne['sous_total']))
        y -= 16

    p.line(50, y - 5, 550, y - 5)
    y -= 28
    p.setFont("Helvetica-Bold", 12)
    p.drawString(360, y, "TOTAL :")
    p.drawRightString(550, y, fm(vente['total']))
    p.setFont("Helvetica", 9)
    p.drawString(50, 40, "Merci pour votre confiance.")

    p.showPage()
    p.save()
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name=f"facture_{vente['id']}.pdf",
                     mimetype='application/pdf')


@app.route('/ventes')
@login_required
def ventes():
    date_filtre = request.args.get('date', '')
    return render_template('ventes.html', ventes=db.recuperer_ventes(date_filtre), date_filtre=date_filtre)


# ----------------------------------------------------------------------
# DÉPENSES
# ----------------------------------------------------------------------
@app.route('/depenses')
@login_required
def depenses():
    depenses_list = db.recuperer_depenses()
    total_depenses = sum(d['montant'] for d in depenses_list)
    return render_template('depenses.html', depenses=depenses_list, total_depenses=total_depenses)


@app.route('/depenses/ajouter', methods=['POST'])
@login_required
def ajouter_depense():
    description = request.form.get('description', '').strip()
    categorie = request.form.get('categorie', '').strip()
    try:
        montant = float(request.form.get('montant', 0))
    except ValueError:
        flash('Montant invalide.', 'error')
        return redirect(url_for('depenses'))
    if description and categorie and montant > 0:
        db.ajouter_depense(description, categorie, montant)
        flash('Dépense enregistrée avec succès.', 'success')
    else:
        flash('Veuillez remplir tous les champs avec un montant positif.', 'error')
    return redirect(url_for('depenses'))


@app.route('/depenses/supprimer/<int:depense_id>', methods=['POST'])
@login_required
@admin_required
def supprimer_depense(depense_id):
    db.supprimer_depense(depense_id)
    flash('Dépense supprimée avec succès.', 'success')
    return redirect(url_for('depenses'))


# ----------------------------------------------------------------------
# RAPPORTS (6 derniers mois, même source que le dashboard)
# ----------------------------------------------------------------------
@app.route('/rapports')
@login_required
def rapports():
    rapport = db.recuperer_rapport_mensuel(6)  # ordre chronologique
    total_ca = sum(r['ca'] for r in rapport)
    total_benefice = sum(r['benefice'] for r in rapport)
    mois_actifs = [r for r in rapport if r['ca'] > 0]
    meilleur_mois = max(mois_actifs, key=lambda r: r['ca'])['mois'] if mois_actifs else None
    marge = round(total_benefice / total_ca * 100, 1) if total_ca > 0 else 0.0
    return render_template('rapports.html', rapport=rapport,
                           total_benefice=total_benefice, meilleur_mois=meilleur_mois, marge=marge)


@app.route('/rapports/excel')
@login_required
def exporter_excel():
    param = lire_parametres()
    cles = param.keys() if param else []
    taux_ar = param['taux_ar'] if 'taux_ar' in cles and param['taux_ar'] else 5021.0
    devise = session.get('devise', param['devise_defaut'] if 'devise_defaut' in cles else 'Ar')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rapport Financier"
    entetes = ['Mois', f"Chiffre d'Affaires ({devise})", f'Dépenses ({devise})',
               f'Bénéfice Net ({devise})', 'Ventes Réalisées']
    ws.append(entetes)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F2BD9")
        cell.alignment = Alignment(horizontal="center")

    for r in db.recuperer_rapport_mensuel(6):
        ws.append([r['mois'],
                   round(convertir_montant(r['ca'], devise, taux_ar), 2),
                   round(convertir_montant(r['depenses'], devise, taux_ar), 2),
                   round(convertir_montant(r['benefice'], devise, taux_ar), 2),
                   r['nb_ventes']])

    for col in ('B', 'C', 'D'):
        ws.column_dimensions[col].width = 24
        for cell in ws[col][1:]:
            cell.number_format = '#,##0' if devise == 'Ar' else '#,##0.00'
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['E'].width = 18

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name='rapport_financier.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ----------------------------------------------------------------------
# PAGE 404
# ----------------------------------------------------------------------
@app.errorhandler(404)
def page_not_found(e):
    return render_template('404.html'), 404


if __name__ == '__main__':
    # debug=True uniquement en local. Désactiver avant toute mise en ligne.
    app.run(debug=True)